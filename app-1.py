"""
Bijli Tools server (Phase 3B).

Endpoints (sab POST multipart, field "file"; Authorization: Bearer <API_TOKEN>):
  GET  /health
  POST /v1/pdf-to-docx
  POST /v1/pdf-to-xlsx      form: text_fallback=1
  POST /v1/pdf-to-pptx      form: mode=image|text
  POST /v1/office-to-pdf
  POST /v1/pdf-protect      form: password
  POST /v1/pdf-unprotect    form: password

Files sirf request ke dauran ek temp folder me rehti hain aur jawab dene ke turant baad delete ho jati hain.
"""
import os, re, time, shutil, signal, tempfile, subprocess, threading, hmac, logging, contextlib, collections, secrets

from flask import Flask, request, jsonify, Response, g
from werkzeug.exceptions import RequestEntityTooLarge, HTTPException
import pikepdf
import pdfplumber
from docx import Document
from docx.shared import Pt, Inches
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Emu, Inches as PInches, Pt as PPt

API_TOKEN = os.environ.get("API_TOKEN", "")
ALLOWED = os.environ.get("ALLOWED_ORIGINS", "*").strip()
ALLOWED_LIST = [o.strip() for o in ALLOWED.split(",") if o.strip()]
MAX_MB = int(os.environ.get("MAX_MB", "25"))
MAX_PAGES = int(os.environ.get("MAX_PAGES", "100"))
CONCURRENCY = int(os.environ.get("CONCURRENCY", "2"))
RATE_PER_MIN = int(os.environ.get("RATE_PER_MIN", "30"))
TRUST_PROXY = os.environ.get("TRUST_PROXY", "") == "1"
SOFFICE = os.environ.get("SOFFICE", "soffice")
OFFICE_TIMEOUT = int(os.environ.get("OFFICE_TIMEOUT", "120"))

if len(API_TOKEN) < 16:
    raise SystemExit("API_TOKEN set karo (kam se kam 16 characters). Banane ke liye: python -c \"import secrets;print(secrets.token_urlsafe(32))\"")

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(message)s")
log = logging.getLogger("bijli")

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_MB * 1024 * 1024
SEM = threading.BoundedSemaphore(CONCURRENCY)
_hits = collections.defaultdict(collections.deque)
_hits_lock = threading.Lock()


class ApiError(Exception):
    def __init__(self, msg, code=400):
        super().__init__(msg)
        self.msg, self.code = msg, code


# ---------- gate: CORS, auth, rate limit ----------
@app.after_request
def _after(resp):
    o = request.headers.get("Origin")
    if o and (ALLOWED == "*" or o in ALLOWED_LIST):
        resp.headers["Access-Control-Allow-Origin"] = "*" if ALLOWED == "*" else o
        resp.headers["Vary"] = "Origin"
        resp.headers["Access-Control-Allow-Headers"] = "Authorization, Content-Type"
        resp.headers["Access-Control-Allow-Methods"] = "GET, POST, OPTIONS"
        resp.headers["Access-Control-Expose-Headers"] = "Content-Disposition"
        resp.headers["Access-Control-Max-Age"] = "600"
    resp.headers["Cache-Control"] = "no-store"
    resp.headers["X-Content-Type-Options"] = "nosniff"
    log.info("%s %s %s in=%sB %.1fs", request.method, request.path, resp.status_code,
             request.content_length or 0, time.time() - getattr(g, "t0", time.time()))
    return resp


def _ip():
    if TRUST_PROXY:
        xf = request.headers.get("X-Forwarded-For", "")
        if xf:
            return xf.split(",")[0].strip()
    return request.remote_addr or "?"


@app.before_request
def _gate():
    g.t0 = time.time()
    if request.method == "OPTIONS":
        return ("", 204)
    ip, now = _ip(), time.time()
    with _hits_lock:
        q = _hits[ip]
        while q and now - q[0] > 60:
            q.popleft()
        if len(q) >= RATE_PER_MIN:
            raise ApiError("Bahut zyada requests. Thodi der ruko.", 429)
        q.append(now)
        if len(_hits) > 5000:  # memory guard
            for k in [k for k, v in _hits.items() if not v or now - v[-1] > 60]:
                _hits.pop(k, None)
    a = request.headers.get("Authorization", "")
    tok = a[7:] if a.startswith("Bearer ") else ""
    if not hmac.compare_digest(tok.encode(), API_TOKEN.encode()):
        raise ApiError("Token galat hai.", 401)


@app.errorhandler(ApiError)
def _api_err(e):
    return jsonify(error=e.msg), e.code


@app.errorhandler(RequestEntityTooLarge)
def _too_big(e):
    return jsonify(error="File bahut badi hai (limit %d MB)." % MAX_MB), 413


@app.errorhandler(HTTPException)
def _http_err(e):
    return jsonify(error=e.description or e.name), e.code


@app.errorhandler(Exception)
def _any_err(e):
    log.exception("unhandled")
    return jsonify(error="Server me dikkat aayi. File kharab ho sakti hai."), 500


# ---------- helpers ----------
@contextlib.contextmanager
def job():
    if not SEM.acquire(timeout=30):
        raise ApiError("Server abhi busy hai, thodi der baad try karo.", 503)
    d = tempfile.mkdtemp(prefix="bt_")
    try:
        yield d
    finally:
        shutil.rmtree(d, ignore_errors=True)
        SEM.release()


def get_upload(d, exts):
    f = request.files.get("file")
    if not f or not f.filename:
        raise ApiError("File nahi mili.", 400)
    ext = f.filename.rsplit(".", 1)[-1].lower() if "." in f.filename else ""
    if ext not in exts:
        raise ApiError("Ye file type support nahi hai: ." + ext, 415)
    p = os.path.join(d, "in." + ext)  # user ka naam kabhi path me nahi jata
    f.save(p)
    if os.path.getsize(p) == 0:
        raise ApiError("File khali hai.", 400)
    return p


def check_pdf(p):
    with open(p, "rb") as fh:
        if b"%PDF" not in fh.read(1024):
            raise ApiError("Ye valid PDF nahi hai.", 400)
    try:
        with pikepdf.open(p) as pdf:
            n = len(pdf.pages)
    except pikepdf.PasswordError:
        raise ApiError('PDF password-protected hai. Pehle "Password hatao" use karo.', 422)
    except pikepdf.PdfError:
        raise ApiError("PDF kharab hai ya khul nahi rahi.", 400)
    if n > MAX_PAGES:
        raise ApiError("PDF me %d pages hain, limit %d hai." % (n, MAX_PAGES), 413)
    return n


def send(path, mime, name):
    with open(path, "rb") as fh:
        data = fh.read()
    return Response(data, mimetype=mime, headers={"Content-Disposition": 'attachment; filename="%s"' % name})


_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff]")
clean = lambda s: _ILLEGAL.sub("", "" if s is None else str(s))
_NUM1 = re.compile(r"^-?(0|[1-9]\d{0,14})(\.\d+)?$")
_NUM2 = re.compile(r"^-?[1-9]\d{0,2}(,\d{3})+(\.\d+)?$")


def typed(v):
    s = clean(v).strip()
    if _NUM1.match(s) and len(s.replace("-", "").replace(".", "")) <= 15:
        return float(s) if "." in s else int(s)
    if _NUM2.match(s):
        t = s.replace(",", "")
        return float(t) if "." in t else int(t)
    return clean(v)


# ---------- PDF -> DOCX ----------
def _lines(words):
    words = sorted(words, key=lambda w: (w["top"], w["x0"]))
    out = []
    for w in words:
        if out and abs(w["top"] - out[-1]["top"]) <= max(2.0, 0.5 * w["size"]):
            out[-1]["words"].append(w)
            out[-1]["bottom"] = max(out[-1]["bottom"], w["bottom"])
        else:
            out.append({"top": w["top"], "bottom": w["bottom"], "words": [w]})
    for ln in out:
        ln["words"].sort(key=lambda w: w["x0"])
        ln["size"] = sum(w["size"] for w in ln["words"]) / len(ln["words"])
    return out


_BULLET = re.compile(r"^(?:[•●▪■◦‣\-–*]|\d{1,3}[.)])\s")


def _paragraphs(lines):
    paras, cur = [], None
    for ln in lines:
        first = clean(ln["words"][0]["text"]) + " "
        new = (
            cur is None
            or ln["top"] - cur["bottom"] > 0.5 * ln["size"]
            or abs(ln["size"] - cur["size"]) > 1.5
            or bool(_BULLET.match(first + "x"))
        )
        if new:
            cur = {"top": ln["top"], "bottom": ln["bottom"], "size": ln["size"], "words": list(ln["words"])}
            paras.append(cur)
        else:
            cur["words"].extend(ln["words"])
            cur["bottom"] = ln["bottom"]
    return paras


def _in_box(w, b):
    cx, cy = (w["x0"] + w["x1"]) / 2, (w["top"] + w["bottom"]) / 2
    return b[0] <= cx <= b[2] and b[1] <= cy <= b[3]


def pdf_to_docx(src, dst):
    pages = []
    with pdfplumber.open(src) as pdf:
        for pg in pdf.pages:
            tabs = pg.find_tables()
            boxes = [t.bbox for t in tabs]
            words = [
                w for w in pg.extract_words(extra_attrs=["size", "fontname"], keep_blank_chars=False)
                if not any(_in_box(w, b) for b in boxes)
            ]
            tdata = []
            for t in tabs:
                rows = t.extract() or []
                rows = [[clean(c) for c in r] for r in rows if r]
                if rows and any(any(c.strip() for c in r) for r in rows):
                    tdata.append((t.bbox[1], rows))
            pages.append({"w": float(pg.width), "h": float(pg.height), "words": words, "tables": tdata})
    weight = collections.Counter()
    for p in pages:
        for w in p["words"]:
            weight[round(w["size"] * 2) / 2] += len(w["text"])
    chars = sum(weight.values())
    if chars < 20 and not any(p["tables"] for p in pages):
        raise ApiError("Is PDF me selectable text nahi mila (scan ki hui lagti hai). Pehle OCR karo.", 422)
    body = weight.most_common(1)[0][0] if weight else 11

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Pt(pages[0]["w"]), Pt(pages[0]["h"])
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, m, Inches(0.8))
    doc.styles["Normal"].font.size = Pt(max(8, min(14, body)))
    for i, p in enumerate(pages):
        if i:
            doc.add_page_break()
        blocks = [(pa["top"], "p", pa) for pa in _paragraphs(_lines(p["words"]))]
        blocks += [(top, "t", rows) for top, rows in p["tables"]]
        blocks.sort(key=lambda b: b[0])
        for _, kind, data in blocks:
            if kind == "t":
                ncol = max(len(r) for r in data)
                tb = doc.add_table(rows=len(data), cols=ncol)
                tb.style = "Table Grid"
                for ri, r in enumerate(data):
                    for ci in range(ncol):
                        tb.cell(ri, ci).text = r[ci] if ci < len(r) else ""
                doc.add_paragraph()
                continue
            words = data["words"]
            size = sum(w["size"] for w in words) / len(words)
            if size >= body * 1.25 and len(words) <= 25:
                lvl = 1 if size >= body * 1.7 else 2 if size >= body * 1.4 else 3
                doc.add_heading(clean(" ".join(w["text"] for w in words)), level=lvl)
                continue
            para = doc.add_paragraph()
            runs = []
            for w in words:
                fn = (w.get("fontname") or "").lower()
                st = ("bold" in fn or "black" in fn or "heavy" in fn, "italic" in fn or "oblique" in fn)
                if runs and runs[-1][1] == st:
                    runs[-1][0] += " " + clean(w["text"])
                else:
                    if runs:
                        runs[-1][0] += " "
                    runs.append([clean(w["text"]), st])
            for text, (b, it) in runs:
                r = para.add_run(text)
                r.bold, r.italic = b or None, it or None
                r.font.size = Pt(max(6, round(size * 2) / 2))
    doc.save(dst)
    return chars


# ---------- PDF -> XLSX ----------
def pdf_to_xlsx(src, dst, text_fallback):
    wb = Workbook()
    wb.remove(wb.active)
    with pdfplumber.open(src) as pdf:
        for pi, pg in enumerate(pdf.pages, 1):
            tabs = [t for t in (pg.extract_tables() or []) if t]
            made = 0
            for ti, t in enumerate(tabs, 1):
                rows = [r for r in t if r and any((c or "").strip() for c in r)]
                if not rows:
                    continue
                ws = wb.create_sheet("P%d_T%d" % (pi, ti))
                for r in rows:
                    ws.append([typed(c) if c is not None else None for c in r])
                made += 1
            if not made and text_fallback:
                txt = pg.extract_text(layout=True) or ""
                lines = [ln for ln in txt.splitlines() if ln.strip()]
                if lines:
                    ws = wb.create_sheet("P%d_Text" % pi)
                    for ln in lines:
                        ws.append([typed(c) for c in re.split(r"\s{2,}", ln.strip())])
    if not wb.sheetnames:
        raise ApiError("Koi table nahi mili. 'Table na mile to text bhi lo' option try karo.", 422)
    wb.save(dst)
    return len(wb.sheetnames)


# ---------- PDF -> PPTX ----------
def _fit(prs, pic_w, pic_h):
    sw, sh = prs.slide_width, prs.slide_height
    s = min(sw / pic_w, sh / pic_h)
    w, h = int(pic_w * s), int(pic_h * s)
    return (sw - w) // 2, (sh - h) // 2, w, h


def pdf_to_pptx(src, dst, d, mode, n):
    from PIL import Image
    prs = Presentation()
    blank = prs.slide_layouts[6]
    if mode == "image":
        if not shutil.which("pdftoppm"):
            raise ApiError("Server par pdftoppm (poppler) install nahi hai.", 501)
        pre = os.path.join(d, "pg")
        try:
            subprocess.run(["pdftoppm", "-r", "110", "-jpeg", "-f", "1", "-l", str(n), src, pre],
                           check=True, timeout=120, capture_output=True)
        except (subprocess.TimeoutExpired, subprocess.CalledProcessError):
            raise ApiError("Pages ko image banane me dikkat aayi.", 500)
        imgs = sorted(f for f in os.listdir(d) if f.startswith("pg") and f.endswith(".jpg"))
        if not imgs:
            raise ApiError("Pages render nahi hue.", 500)
        with Image.open(os.path.join(d, imgs[0])) as im0:
            ar = im0.height / im0.width
        prs.slide_width = Emu(12192000)
        prs.slide_height = Emu(int(12192000 * min(max(ar, .4), 1.5)))
        for f in imgs:
            p = os.path.join(d, f)
            with Image.open(p) as im:
                x, y, w, h = _fit(prs, im.width, im.height)
            prs.slides.add_slide(blank).shapes.add_picture(p, x, y, w, h)
    else:
        prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
        count = 0
        with pdfplumber.open(src) as pdf:
            for pg in pdf.pages:
                txt = clean(pg.extract_text() or "").strip()
                count += len(txt)
                s = prs.slides.add_slide(blank)
                tb = s.shapes.add_textbox(PInches(.5), PInches(.4), PInches(12.33), PInches(6.7))
                tf = tb.text_frame
                tf.word_wrap = True
                size = 16 if len(txt) < 700 else 12 if len(txt) < 1600 else 9
                for i, line in enumerate(txt.split("\n") if txt else [""]):
                    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    para.text = line
                    for r in para.runs:
                        r.font.size = PPt(size)
        if count < 20:
            raise ApiError("Is PDF me selectable text nahi mila (scan). 'Image slides' mode try karo.", 422)
    prs.save(dst)


# ---------- Office -> PDF ----------
OFFICE_EXT = {"doc", "docx", "xls", "xlsx", "ppt", "pptx", "odt", "ods", "odp", "rtf", "txt", "csv"}


def office_to_pdf(src, d):
    if not shutil.which(SOFFICE):
        raise ApiError("Server par LibreOffice install nahi hai.", 501)
    prof = tempfile.mkdtemp(prefix="lo_")
    try:
        cmd = [SOFFICE, "-env:UserInstallation=file://" + prof, "--headless", "--norestore", "--nolockcheck",
               "--convert-to", "pdf", "--outdir", d, src]
        pr = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, start_new_session=True)
        try:
            pr.communicate(timeout=OFFICE_TIMEOUT)
        except subprocess.TimeoutExpired:
            os.killpg(pr.pid, signal.SIGKILL)
            pr.communicate()
            raise ApiError("Conversion me bahut der lag gayi (file bhari ya password wali ho sakti hai).", 504)
    finally:
        shutil.rmtree(prof, ignore_errors=True)
    out = os.path.splitext(src)[0] + ".pdf"
    if not os.path.exists(out) or os.path.getsize(out) == 0:
        raise ApiError("File convert nahi ho payi. Password wali ya kharab file ho sakti hai.", 422)
    return out


# ---------- routes ----------
@app.get("/health")
def health():
    return jsonify(ok=True, version=1, max_mb=MAX_MB, max_pages=MAX_PAGES,
                   libreoffice=bool(shutil.which(SOFFICE)), poppler=bool(shutil.which("pdftoppm")))


@app.post("/v1/pdf-to-docx")
def r_docx():
    with job() as d:
        src = get_upload(d, {"pdf"})
        check_pdf(src)
        dst = os.path.join(d, "out.docx")
        pdf_to_docx(src, dst)
        return send(dst, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "converted.docx")


@app.post("/v1/pdf-to-xlsx")
def r_xlsx():
    with job() as d:
        src = get_upload(d, {"pdf"})
        check_pdf(src)
        dst = os.path.join(d, "out.xlsx")
        pdf_to_xlsx(src, dst, request.form.get("text_fallback") == "1")
        return send(dst, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "converted.xlsx")


@app.post("/v1/pdf-to-pptx")
def r_pptx():
    mode = request.form.get("mode", "image")
    if mode not in ("image", "text"):
        raise ApiError("mode galat hai.", 400)
    with job() as d:
        src = get_upload(d, {"pdf"})
        n = check_pdf(src)
        dst = os.path.join(d, "out.pptx")
        pdf_to_pptx(src, dst, d, mode, min(n, 60 if mode == "image" else MAX_PAGES))
        return send(dst, "application/vnd.openxmlformats-officedocument.presentationml.presentation", "converted.pptx")


@app.post("/v1/office-to-pdf")
def r_office():
    with job() as d:
        src = get_upload(d, OFFICE_EXT)
        out = office_to_pdf(src, d)
        return send(out, "application/pdf", "converted.pdf")


@app.post("/v1/pdf-protect")
def r_protect():
    pw = request.form.get("password", "")
    if not (1 <= len(pw.encode()) <= 100):
        raise ApiError("Password 1 se 100 characters ka rakho.", 400)
    with job() as d:
        src = get_upload(d, {"pdf"})
        check_pdf(src)
        dst = os.path.join(d, "out.pdf")
        with pikepdf.open(src) as pdf:
            pdf.save(dst, encryption=pikepdf.Encryption(user=pw, owner=secrets.token_urlsafe(24), R=6,
                                                        allow=pikepdf.Permissions()))
        return send(dst, "application/pdf", "protected.pdf")


@app.post("/v1/pdf-unprotect")
def r_unprotect():
    pw = request.form.get("password", "")
    with job() as d:
        src = get_upload(d, {"pdf"})
        with open(src, "rb") as fh:
            if b"%PDF" not in fh.read(1024):
                raise ApiError("Ye valid PDF nahi hai.", 400)
        try:
            with pikepdf.open(src, password=pw) as pdf:
                if not pdf.is_encrypted:
                    raise ApiError("Is PDF me password hai hi nahi.", 400)
                dst = os.path.join(d, "out.pdf")
                pdf.save(dst)  # default: encryption hat jati hai
        except pikepdf.PasswordError:
            raise ApiError("Password galat hai.", 401)
        except pikepdf.PdfError:
            raise ApiError("PDF kharab hai ya khul nahi rahi.", 400)
        return send(dst, "application/pdf", "unlocked.pdf")


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", "8080")), threaded=True)
